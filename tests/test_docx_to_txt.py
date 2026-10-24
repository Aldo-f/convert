from click.testing import CliRunner
from docx import Document

from convert.cli import cli


def test_docx_to_txt(tmp_path):
    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_paragraph("Hello world")
    doc.add_paragraph("Second paragraph")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "2"
    doc.save(src)

    out = tmp_path / "out.txt"
    r = CliRunner().invoke(cli, ["convert", str(src), str(out)])
    assert r.exit_code == 0

    text = out.read_text(encoding="utf-8")
    assert "Hello world" in text
    assert "Second paragraph" in text
    assert "A | B" in text
    assert "1 | 2" in text
