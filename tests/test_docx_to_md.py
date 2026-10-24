from click.testing import CliRunner
from docx import Document

from convert.cli import cli


def test_docx_to_md(tmp_path):
    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_heading("Title", level=1)
    doc.add_heading("Sub", level=2)
    doc.add_paragraph("Body text", style="List Bullet")
    doc.add_paragraph("Quote text", style="Intense Quote")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    doc.save(src)

    out = tmp_path / "out.md"
    r = CliRunner().invoke(cli, ["convert", str(src), str(out)])
    assert r.exit_code == 0

    md = out.read_text(encoding="utf-8")
    assert "# Title" in md
    assert "## Sub" in md
    assert "- Body text" in md
    assert "> Quote text" in md
    assert "| A | B |" in md
