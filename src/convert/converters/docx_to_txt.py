from docx import Document

from convert.models import ConversionRequest, ConversionResult
from convert.registry import register


@register("docx", "txt")
def convert(req: ConversionRequest) -> ConversionResult:
    doc = Document(req.source_path)
    out_lines = []
    for p in doc.paragraphs:
        if p.text:
            out_lines.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            out_lines.append(" | ".join(cells))
        out_lines.append("")
    req.target_path.write_text("\n".join(out_lines), encoding="utf-8")
    return ConversionResult(req.target_path, "docx→txt done")
