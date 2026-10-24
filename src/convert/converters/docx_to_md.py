from docx import Document

from convert.models import ConversionRequest, ConversionResult
from convert.registry import register


@register("docx", "md")
def convert(req: ConversionRequest) -> ConversionResult:
    doc = Document(req.source_path)
    out_lines = []
    for p in doc.paragraphs:
        style = p.style.name
        text = p.text
        if not text:
            continue
        if style == "Heading 1":
            out_lines.append(f"# {text}")
        elif style == "Heading 2":
            out_lines.append(f"## {text}")
        elif style == "Heading 3":
            out_lines.append(f"### {text}")
        elif style == "Heading 4":
            out_lines.append(f"#### {text}")
        elif style == "Heading 5":
            out_lines.append(f"##### {text}")
        elif style == "Heading 6":
            out_lines.append(f"###### {text}")
        elif style == "List Bullet":
            out_lines.append(f"- {text}")
        elif style == "List Number":
            out_lines.append(f"1. {text}")
        elif style == "Intense Quote":
            out_lines.append(f"> {text}")
        else:
            out_lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.replace("\n", " ") for cell in row.cells]
            out_lines.append("| " + " | ".join(cells) + " |")
        out_lines.append("")
    req.target_path.write_text("\n".join(out_lines), encoding="utf-8")
    return ConversionResult(req.target_path, "docx→md done")
