import re

from docx import Document

from convert.models import ConversionRequest, ConversionResult
from convert.registry import register


def parse_inline(text):
    pattern = r"(\*\*[^*]+\*\*|\*[^*]+\*|__[^_]+__|_[^_]+_)"
    parts = re.split(pattern, text)
    result = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            result.append((part[2:-2], True, False))
        elif part.startswith("__") and part.endswith("__"):
            result.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*"):
            result.append((part[1:-1], False, True))
        elif part.startswith("_") and part.endswith("_"):
            result.append((part[1:-1], False, True))
        else:
            result.append((part, False, False))
    return result


def add_formatted_paragraph(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    for txt, bold, italic in parse_inline(text):
        run = p.add_run(txt)
        if bold:
            run.font.bold = True
        if italic:
            run.font.italic = True
    return p


@register("md", "docx")
def convert(req: ConversionRequest) -> ConversionResult:
    with open(req.source_path, encoding="utf-8") as f:
        content = f.read()
    lines = content.splitlines()
    doc = Document()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith("##### "):
            doc.add_heading(stripped[6:], level=5)
        elif stripped.startswith("###### "):
            doc.add_heading(stripped[7:], level=6)
        elif stripped == "" or stripped in ("---", "***", "___"):
            doc.add_paragraph("")
        elif stripped.startswith("|"):
            # Table
            table_lines = []
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                s = lines[j].strip()
                # Skip separator row with dashes
                if s.replace("-", "").replace("|", "").strip() == "":
                    j += 1
                    continue
                table_lines.append(s)
                j += 1
            if table_lines:
                header = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                data = []
                for row in table_lines[1:]:
                    data.append([c.strip() for c in row.split("|")[1:-1]])
                table = doc.add_table(rows=1 + len(data), cols=len(header))
                table.style = "Light Grid Accent 1"
                for idx, h in enumerate(header):
                    cell = table.rows[0].cells[idx]
                    cell.text = h
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.bold = True
                for r_idx, row_data in enumerate(data):
                    for c_idx, text in enumerate(row_data):
                        if c_idx < len(table.rows[r_idx + 1].cells):
                            table.rows[r_idx + 1].cells[c_idx].text = text
            i = j - 1
        elif stripped.startswith("- ") or stripped.startswith("* "):
            add_formatted_paragraph(doc, stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s+", stripped):
            add_formatted_paragraph(doc, re.sub(r"^\d+\.\s+", "", stripped), style="List Number")
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.style = doc.styles["Intense Quote"]
            for txt, b, italic in parse_inline(stripped[2:]):
                run = p.add_run(txt)
                if b:
                    run.font.bold = True
                if italic:
                    run.font.italic = True
        else:
            # Remove double-pipe table artifacts if any
            cleaned = stripped
            if cleaned.startswith("|| "):
                cleaned = cleaned[3:]
            add_formatted_paragraph(doc, cleaned)
        i += 1
    doc.save(req.target_path)
    return ConversionResult(req.target_path, "md->docx done")
